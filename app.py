"""DocShift - convert documents and edit PDFs. Everything runs on your own computer."""
import os, sys, io, re, json, uuid, shutil, tempfile, subprocess, zipfile, time, threading, base64, atexit, webbrowser, statistics
from pathlib import Path
from flask import Flask, request, send_file, jsonify, send_from_directory

HERE = Path(__file__).resolve().parent
PORT = 8767
TMP = Path(tempfile.mkdtemp(prefix="docshift_"))
atexit.register(lambda: shutil.rmtree(TMP, ignore_errors=True))
app = Flask(__name__, static_folder=str(HERE / "static"), static_url_path="/s")
last_ping = [0.0]
NOWIN = getattr(subprocess, "CREATE_NO_WINDOW", 0)
IMG = {"png", "jpg", "jpeg", "webp", "bmp", "tiff", "tif", "gif"}
LO_FMT = {"pdf": "pdf", "docx": "docx", "odt": "odt", "rtf": "rtf", "txt": "txt:Text", "html": "html", "csv": "csv", "xlsx": "xlsx", "ods": "ods", "odp": "odp", "pptx": "pptx"}
TARGETS = {
    "pdf": ["docx", "xlsx", "pptx", "txt", "png", "jpg"],
    **{e: ["pdf", "docx", "odt", "rtf", "txt", "html"] for e in ("docx", "doc", "odt", "rtf")},
    "txt": ["pdf", "docx"], "md": ["pdf", "docx"], "html": ["pdf", "docx"], "htm": ["pdf", "docx"],
    **{e: ["pdf", "csv", "ods"] for e in ("xlsx", "xls", "ods")}, "csv": ["pdf", "xlsx"],
    **{e: ["pdf", "odp"] for e in ("pptx", "ppt", "odp")},
    **{e: ["pdf", "png", "jpg", "webp"] for e in IMG},
}

def ext(n): return n.rsplit(".", 1)[-1].lower() if "." in n else ""
def fail(msg, code=400): return jsonify(error=str(msg)), code
def workdir():
    d = TMP / uuid.uuid4().hex[:8]; d.mkdir(); return d
def save_upload(f, d):
    p = Path(d) / re.sub(r"[^\w.\- ]", "_", f.filename or "file"); f.save(p); return p
def send(path, name=None, notice=""):
    r = send_file(str(path), as_attachment=True, download_name=name or Path(path).name)
    r.headers["X-Notice"] = notice; return r
def parse_pages(spec, n):
    spec = (spec or "").strip().lower()
    if spec in ("", "all"): return list(range(n))
    out = []
    for t in re.split(r"[,\s]+", spec):
        if not t: continue
        m = re.fullmatch(r"(\d*)-(\d*)", t)
        if m: a, b = int(m[1] or 1), int(m[2] or n)
        elif t.isdigit(): a = b = int(t)
        else: raise ValueError(f"Can't understand '{t}'. Use something like 1-3, 5")
        if a < 1 or b > n or a > b: raise ValueError(f"Pages must be between 1 and {n}")
        out += range(a - 1, b)
    return out

# ---------- office conversions (LibreOffice) ----------
def soffice():
    for p in (shutil.which("soffice"), shutil.which("libreoffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if p and os.path.exists(p): return p
def lo_convert(src, fmt, outdir, infilter=None):
    so = soffice()
    if not so: raise RuntimeError("LibreOffice is needed for this conversion. Install it free from libreoffice.org, then try again.")
    prof = Path(tempfile.mkdtemp(dir=TMP))
    cmd = [so, f"-env:UserInstallation={prof.as_uri()}", "--headless", *([f"--infilter={infilter}"] if infilter else []), "--convert-to", fmt, "--outdir", str(outdir), str(src)]
    subprocess.run(cmd, capture_output=True, timeout=300, creationflags=NOWIN)
    hits = list(Path(outdir).glob(Path(src).stem + "." + fmt.split(":")[0]))
    if not hits: raise RuntimeError("The conversion did not produce a file. The document may be damaged or protected.")
    return hits[0]

# ---------- PDF -> other ----------
def pdf_images(src, scale=2):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(src))
    for i in range(len(pdf)): yield pdf[i].render(scale=scale).to_pil().convert("RGB")
def num(s):
    if isinstance(s, str):
        t = s.replace(",", "").strip()
        try: return float(t) if "." in t else int(t)
        except ValueError: return s
    return s if s is not None else ""
def pdf_to_xlsx(src, out):
    import pdfplumber
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active); found = False
    with pdfplumber.open(str(src)) as pdf:
        for n, pg in enumerate(pdf.pages, 1):
            for k, t in enumerate(pg.extract_tables(), 1):
                ws = wb.create_sheet(f"Page{n}_Table{k}"[:31]); found = True
                for r in t: ws.append([num(c) for c in r])
        if not found:
            ws = wb.create_sheet("Text")
            for pg in pdf.pages:
                for line in (pg.extract_text() or "").splitlines(): ws.append([line])
    wb.save(str(out))
    return out, "" if found else "No tables were found, so the text was placed one line per row."
def pdf_to_pptx(src, out):
    from pptx import Presentation
    from pptx.util import Emu
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(src)); w, h = pdf[0].get_size()
    prs = Presentation(); prs.slide_width = Emu(int(w * 12700)); prs.slide_height = Emu(int(h * 12700))
    for im in pdf_images(src, 2):
        b = io.BytesIO(); im.save(b, "PNG"); b.seek(0)
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(b, 0, 0, prs.slide_width, prs.slide_height)
    prs.save(str(out)); return out, "Each page became a slide image. Text on slides is not editable."
def pdf_to_txt(src, out):
    import pdfplumber
    with pdfplumber.open(str(src)) as pdf: txt = "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    Path(out).write_text(txt, encoding="utf-8")
    return out, "" if txt.strip() else "No text found. This looks like a scanned PDF."
def pdf_to_pics(src, tgt, work):
    fmt = "PNG" if tgt == "png" else "JPEG"; stem = Path(src).stem; files = []
    for i, im in enumerate(pdf_images(src, 2), 1):
        p = work / f"{stem}-page{i}.{tgt}"; im.save(p, fmt, **({"quality": 92} if fmt == "JPEG" else {})); files.append(p)
    if len(files) == 1: return files[0], ""
    z = work / f"{stem}-images.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for f in files: zf.write(f, f.name)
    return z, f"{len(files)} images saved in one zip file."
def pdf_to_docx(src, out):
    try:
        from pdf2docx import Converter
        cv = Converter(str(src)); cv.convert(str(out)); cv.close()
        return out, ""
    except Exception:
        pass
    return simple_pdf_to_docx(src, out), "Converted with the built-in engine. Complex layouts may need touch-ups."
def simple_pdf_to_docx(src, out):
    import pdfplumber
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    with pdfplumber.open(str(src)) as pdf:
        p0 = pdf.pages[0]; sec = doc.sections[0]
        sec.page_width, sec.page_height = Inches(p0.width / 72), Inches(p0.height / 72)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(.7)
        sizes = [c["size"] for pg in pdf.pages[:20] for c in pg.chars]
        body = statistics.median(sizes) if sizes else 11
        pics = list(pdf_images(src, 1.6)) if not sizes else []
        for n, pg in enumerate(pdf.pages):
            if n: doc.add_page_break()
            if not sizes:
                b = io.BytesIO(); pics[n].save(b, "PNG"); b.seek(0)
                doc.add_picture(b, width=sec.page_width - sec.left_margin - sec.right_margin); continue
            tb = [t.bbox for t in pg.find_tables()]
            inside = lambda l: any(l["x0"] >= b[0] - 2 and l["x1"] <= b[2] + 2 and l["top"] >= b[1] - 2 and l["bottom"] <= b[3] + 2 for b in tb)
            lines = [l for l in pg.extract_text_lines(return_chars=True) if not inside(l)]
            paras, cur = [], None
            for l in sorted(lines, key=lambda l: (round(l["top"]), l["x0"])):
                sz = statistics.median(c["size"] for c in l["chars"])
                if cur and 0 <= l["top"] - cur["bottom"] < sz * .7 and abs(sz - cur["size"]) < 1.2: cur["lines"].append(l); cur["bottom"] = l["bottom"]
                else: cur = {"lines": [l], "top": l["top"], "bottom": l["bottom"], "size": sz}; paras.append(cur)
            items = [(p["top"], "p", p) for p in paras] + [(t.bbox[1], "t", t) for t in pg.find_tables()] + [(i["top"], "i", i) for i in pg.images]
            for _, kind, o in sorted(items, key=lambda x: x[0]):
                if kind == "p":
                    text = " ".join(l["text"].strip() for l in o["lines"]).strip()
                    if not text: continue
                    chars = [c for l in o["lines"] for c in l["chars"]]; fn = [c["fontname"].lower() for c in chars]
                    bold = sum(("bold" in f or "black" in f) for f in fn) > len(fn) * .6; ital = sum(("italic" in f or "oblique" in f) for f in fn) > len(fn) * .6
                    if o["size"] >= body * 1.25 and len(text) < 200:
                        para = doc.add_heading(text, 1 if o["size"] >= body * 1.8 else 2 if o["size"] >= body * 1.45 else 3)
                    else:
                        para = doc.add_paragraph(); r = para.add_run(text); r.bold, r.italic = bold, ital
                        if abs(o["size"] - 11) > .8: r.font.size = Pt(round(o["size"], 1))
                    l0 = o["lines"][0]
                    if abs((l0["x0"] + l0["x1"]) / 2 - pg.width / 2) < 12 and (l0["x1"] - l0["x0"]) < pg.width * .7: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif kind == "t":
                    rows = o.extract()
                    if not rows: continue
                    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); t.style = "Table Grid"
                    for i, r in enumerate(rows):
                        for j, c in enumerate(r): t.cell(i, j).text = (c or "").replace("\n", " ")
                else:
                    try:
                        bb = (max(o["x0"], 0), max(o["top"], 0), min(o["x1"], pg.width), min(o["bottom"], pg.height))
                        b = io.BytesIO(); pg.crop(bb).to_image(resolution=150).original.save(b, "PNG"); b.seek(0)
                        doc.add_picture(b, width=Inches(min((bb[2] - bb[0]) / 72, 6.5)))
                    except Exception: pass
    doc.save(str(out)); return Path(out)

def images_to_pdf(paths, out):
    from PIL import Image, ImageOps
    ims = [ImageOps.exif_transpose(Image.open(p)).convert("RGB") for p in paths]
    ims[0].save(str(out), save_all=True, append_images=ims[1:], resolution=150); return out
def images_to_images(paths, tgt, work):
    from PIL import Image, ImageOps
    fmt = {"jpg": "JPEG", "png": "PNG", "webp": "WEBP"}[tgt]; outs = []
    for p in paths:
        im = ImageOps.exif_transpose(Image.open(p)); im = im.convert("RGB") if fmt == "JPEG" else im
        o = work / (p.stem + "-converted." + tgt); im.save(o, fmt); outs.append(o)
    if len(outs) == 1: return outs[0]
    z = work / "converted-images.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for o in outs: zf.write(o, o.name)
    return z

def convert(paths, src, tgt, work):
    p = paths[0]; stem = p.stem
    if src == "pdf":
        if tgt == "docx": return pdf_to_docx(p, work / (stem + ".docx"))
        if tgt == "xlsx": return pdf_to_xlsx(p, work / (stem + ".xlsx"))
        if tgt == "pptx": return pdf_to_pptx(p, work / (stem + ".pptx"))
        if tgt == "txt": return pdf_to_txt(p, work / (stem + ".txt"))
        if tgt in ("png", "jpg"): return pdf_to_pics(p, tgt, work)
        raise ValueError("That conversion isn't supported")
    if src in IMG:
        if tgt == "pdf": return images_to_pdf(paths, work / (stem + ".pdf")), ""
        if tgt in ("png", "jpg", "webp"): return images_to_images(paths, tgt, work), ""
        raise ValueError("That conversion isn't supported")
    if src == "md":
        import markdown
        h = work / (stem + ".html")
        h.write_text("<meta charset='utf-8'>" + markdown.markdown(p.read_text("utf-8", errors="replace"), extensions=["tables", "fenced_code"]), encoding="utf-8")
        p, src = h, "html"
    if tgt not in LO_FMT or tgt not in TARGETS.get(src, []): raise ValueError("That conversion isn't supported")
    html = src in ("html", "htm")
    return lo_convert(p, LO_FMT[tgt], work, "HTML (StarWriter)" if html else None), ""

@app.post("/api/convert")
def api_convert():
    files = request.files.getlist("files"); tgt = request.form.get("target", "").lower()
    if not files: return fail("No file received")
    work = workdir(); paths = [save_upload(f, work) for f in files]
    try: out, notice = convert(paths, ext(paths[0].name), tgt, work)
    except Exception as e: return fail(e, 500)
    return send(out, notice=notice)
@app.get("/api/formats")
def api_formats(): return jsonify(targets=TARGETS, office=bool(soffice()))

# ---------- PDF editor ----------
SESS = {}
@app.post("/api/open")
def api_open():
    import pypdfium2 as pdfium
    f = request.files["file"]; sid = uuid.uuid4().hex[:10]; d = TMP / sid; d.mkdir(); p = d / "doc.pdf"; f.save(p)
    try:
        pdf = pdfium.PdfDocument(str(p)); sizes = [list(pdf[i].get_size()) for i in range(len(pdf))]
    except Exception as e: return fail("This PDF is protected or damaged. Use PDF Tools > Unlock first." if "password" in str(e).lower() else e)
    SESS[sid] = {"path": p, "name": f.filename}; return jsonify(sid=sid, name=f.filename, pages=sizes)
@app.get("/api/page/<sid>/<int:n>")
def api_page(sid, n):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(SESS[sid]["path"])); im = pdf[n].render(scale=min(4, float(request.args.get("s", 1.5)))).to_pil()
    b = io.BytesIO(); im.save(b, "PNG", compress_level=1); b.seek(0); return send_file(b, mimetype="image/png")
def rgb_hex(c):
    try:
        if c is None: return "#111111"
        if isinstance(c, (int, float)): c = (c,)
        if len(c) == 1: v = int(c[0] * 255); return "#%02x%02x%02x" % (v, v, v)
        if len(c) == 3: return "#%02x%02x%02x" % tuple(int(x * 255) for x in c)
        if len(c) == 4: k = c[3]; return "#%02x%02x%02x" % tuple(int(255 * (1 - x) * (1 - k)) for x in c[:3])
    except Exception: pass
    return "#111111"
@app.get("/api/lines/<sid>/<int:n>")
def api_lines(sid, n):
    import pdfplumber
    out = []
    with pdfplumber.open(str(SESS[sid]["path"])) as pdf:
        for l in pdf.pages[n].extract_text_lines(return_chars=True):
            ch = [c for c in l["chars"] if c["text"].strip()] or l["chars"]; fn = ch[0]["fontname"].lower()
            out.append(dict(x=l["x0"], y=l["top"], w=l["x1"] - l["x0"], h=l["bottom"] - l["top"], text=l["text"], size=round(statistics.median(c["size"] for c in ch), 1),
                            bold=("bold" in fn or "black" in fn), italic=("italic" in fn or "oblique" in fn),
                            font="times" if "times" in fn or "serif" in fn and "sans" not in fn else "courier" if "courier" in fn or "mono" in fn else "helv", color=rgb_hex(ch[0].get("non_stroking_color"))))
    return jsonify(out)

_font = {}
def ttf():
    if "n" not in _font:
        _font["n"] = None
        for p in (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\segoeui.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf"):
            if os.path.exists(p):
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont
                pdfmetrics.registerFont(TTFont("DSUni", p)); _font["n"] = "DSUni"; break
    return _font["n"]
def hexrgb(h):
    h = (h or "#000000").lstrip("#"); return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
def rl_font(it):
    base = {"times": "Times", "courier": "Courier"}.get(it.get("font"), "Helvetica"); b, i = it.get("bold"), it.get("italic")
    if base == "Times": return "Times-" + ("BoldItalic" if b and i else "Bold" if b else "Italic" if i else "Roman")
    return base + ("-BoldOblique" if b and i else "-Bold" if b else "-Oblique" if i else "")
def draw_text(c, it, H):
    txt = it.get("text", ""); font = rl_font(it)
    try: txt.encode("cp1252")
    except UnicodeEncodeError: font = ttf() or font
    c.setFont(font, it["size"]); c.setFillColorRGB(*hexrgb(it.get("color"))); c.setFillAlpha(1)
    for k, line in enumerate(txt.split("\n")): c.drawString(it["x"], H - it["y"] - it["size"] * .8 - k * it["size"] * 1.2, line)
def draw_op(c, it, H):
    from reportlab.lib.utils import ImageReader
    t = it["t"]
    if t == "text": draw_text(c, it, H)
    elif t == "rect":
        c.setFillColorRGB(*hexrgb(it.get("fill"))); c.setFillAlpha(float(it.get("alpha", 1))); c.rect(it["x"], H - it["y"] - it["h"], it["w"], it["h"], stroke=0, fill=1)
    elif t == "replace":
        c.setFillColorRGB(1, 1, 1); c.setFillAlpha(1); c.rect(it["x"] - 1.5, H - it["y"] - it["h"] - 1.5, it["w"] + 3, it["h"] + 3, stroke=0, fill=1); draw_text(c, it, H)
    elif t == "draw" and len(it["pts"]) > 1:
        c.setStrokeColorRGB(*hexrgb(it.get("color"))); c.setLineWidth(float(it.get("width", 2))); c.setLineCap(1); c.setLineJoin(1); c.setStrokeAlpha(1)
        p = c.beginPath(); p.moveTo(it["pts"][0][0], H - it["pts"][0][1])
        for x, y in it["pts"][1:]: p.lineTo(x, H - y)
        c.drawPath(p, stroke=1, fill=0)
    elif t == "image":
        img = ImageReader(io.BytesIO(base64.b64decode(it["data"].split(",", 1)[1]))); c.drawImage(img, it["x"], H - it["y"] - it["h"], it["w"], it["h"], mask="auto")
def overlay(w, h, fn):
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader
    b = io.BytesIO(); c = canvas.Canvas(b, pagesize=(w, h)); fn(c); c.save(); b.seek(0); return PdfReader(b).pages[0]
def apply_ops(src, ops, out):
    from pypdf import PdfReader, PdfWriter
    r, w = PdfReader(str(src)), PdfWriter()
    for i, page in enumerate(r.pages):
        items = ops.get(str(i)) or []
        if items:
            if page.rotation: page.transfer_rotation_to_content()
            W, H = float(page.mediabox.width), float(page.mediabox.height)
            page.merge_page(overlay(W, H, lambda c: [draw_op(c, it, H) for it in items]))
        w.add_page(page)
    with open(out, "wb") as f: w.write(f)
@app.post("/api/save/<sid>")
def api_save(sid):
    s = SESS.get(sid)
    if not s: return fail("Session expired. Open the file again.")
    out = TMP / sid / (uuid.uuid4().hex[:6] + ".pdf")
    try: apply_ops(s["path"], request.get_json(force=True), out)
    except Exception as e: return fail(e, 500)
    return send(out, name=Path(s["name"]).stem + "-edited.pdf")

# ---------- PDF tools ----------
@app.post("/api/tool/<name>")
def api_tool(name):
    from pypdf import PdfReader, PdfWriter
    files = request.files.getlist("files"); F = request.form
    if not files: return fail("Choose a PDF first")
    work = workdir(); paths = [save_upload(f, work) for f in files]; stem = paths[0].stem
    try:
        def reader(p):
            r = PdfReader(str(p))
            if r.is_encrypted:
                if not r.decrypt(F.get("password", "")): raise ValueError("This PDF is password protected. Use the Unlock tool, or enter its password.")
            return r
        out = work / (stem + "-" + name + ".pdf"); notice = ""
        if name == "merge":
            w = PdfWriter()
            for p in paths: w.append(reader(p))
            out = work / "merged.pdf"
        else:
            r = reader(paths[0]); n = len(r.pages); w = PdfWriter()
            if name in ("split", "extract"):
                idx = parse_pages(F.get("pages"), n)
                if name == "extract":
                    for i in idx: w.add_page(r.pages[i])
                else:
                    z = work / (stem + "-split.zip")
                    with zipfile.ZipFile(z, "w") as zf:
                        for k in range(0, len(idx), max(1, int(F.get("every", 1) or 1))):
                            ww = PdfWriter()
                            for i in idx[k:k + max(1, int(F.get("every", 1) or 1))]: ww.add_page(r.pages[i])
                            b = io.BytesIO(); ww.write(b); zf.writestr(f"{stem}-{idx[k] + 1}.pdf", b.getvalue())
                    return send(z)
            elif name == "rotate":
                sel = set(parse_pages(F.get("pages"), n)); ang = int(F.get("angle", 90))
                for i, pg in enumerate(r.pages):
                    if i in sel: pg.rotate(ang)
                    w.add_page(pg)
            elif name == "delete":
                drop = set(parse_pages(F.get("pages"), n))
                if len(drop) >= n: raise ValueError("That would delete every page")
                for i, pg in enumerate(r.pages):
                    if i not in drop: w.add_page(pg)
            elif name == "reorder":
                order = parse_pages(F.get("pages"), n); rest = [i for i in range(n) if i not in order]
                for i in order + rest: w.add_page(r.pages[i])
            elif name == "compress":
                for pg in r.pages: w.add_page(pg)
                for pg in w.pages: pg.compress_content_streams()
                w.compress_identical_objects(remove_identicals=True, remove_orphans=True)
                notice = "Compression is lossless, so scanned or image-heavy files may shrink only a little."
            elif name == "protect":
                if not F.get("password"): raise ValueError("Enter a password")
                for pg in r.pages: w.add_page(pg)
                w.encrypt(F["password"])
            elif name == "unlock":
                for pg in r.pages: w.add_page(pg)
            elif name == "watermark":
                text = F.get("text", "CONFIDENTIAL")
                for pg in r.pages:
                    W, H = float(pg.mediabox.width), float(pg.mediabox.height)
                    def fn(c):
                        c.saveState(); c.translate(W / 2, H / 2); c.rotate(45); c.setFillColorRGB(.5, .5, .5); c.setFillAlpha(.22); c.setFont("Helvetica-Bold", min(W, H) / 7)
                        c.drawCentredString(0, 0, text); c.restoreState()
                    pg.merge_page(overlay(W, H, fn)); w.add_page(pg)
            elif name == "numbers":
                for i, pg in enumerate(r.pages, 1):
                    W, H = float(pg.mediabox.width), float(pg.mediabox.height)
                    pg.merge_page(overlay(W, H, lambda c: (c.setFont("Helvetica", 10), c.setFillColorRGB(.2, .2, .2), c.drawCentredString(W / 2, 22, f"{i} / {n}")))); w.add_page(pg)
            else: return fail("Unknown tool", 404)
        with open(out, "wb") as f: w.write(f)
    except Exception as e: return fail(e, 400)
    return send(out, notice=notice)

@app.get("/")
def index(): return send_from_directory(str(HERE / "static"), "index.html")
@app.get("/ping")
def ping(): last_ping[0] = time.time(); return ("", 204)

def open_window(url):
    for exe in (r"%ProgramFiles(x86)%\Microsoft\Edge\Application\msedge.exe", r"%ProgramFiles%\Microsoft\Edge\Application\msedge.exe",
                r"%ProgramFiles%\Google\Chrome\Application\chrome.exe", r"%ProgramFiles(x86)%\Google\Chrome\Application\chrome.exe", r"%LocalAppData%\Google\Chrome\Application\chrome.exe"):
        e = os.path.expandvars(exe)
        if os.path.exists(e): subprocess.Popen([e, f"--app={url}", "--window-size=1320,860"]); return
    webbrowser.open(url)
if __name__ == "__main__":
    import logging, socket
    logging.getLogger("werkzeug").setLevel(logging.ERROR)
    url = f"http://127.0.0.1:{PORT}/"
    with socket.socket() as s:
        if s.connect_ex(("127.0.0.1", PORT)) == 0: open_window(url); sys.exit(0)
    threading.Thread(target=lambda: app.run(host="127.0.0.1", port=PORT, threaded=True), daemon=True).start()
    start = time.time()
    if "--no-open" not in sys.argv: open_window(url)
    while True:
        time.sleep(3)
        if last_ping[0] and time.time() - last_ping[0] > 25: break
        if not last_ping[0] and time.time() - start > 120 and "--no-open" not in sys.argv: break
